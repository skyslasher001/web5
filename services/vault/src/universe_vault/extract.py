"""Text extraction per format and chunking into retrieval units (SPEC §5 pipeline step 3).

Permissive libraries only: pypdf, python-docx, openpyxl, python-pptx (PyMuPDF is AGPL and not used).
"""

import csv
import io
import json
import re
from collections.abc import Iterator
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, ClassVar

from universe_vault.filetypes import PLAIN_TEXT_EXTENSIONS
from universe_vault.metadata import parse_frontmatter

MAX_TEXT_CHARS = 2_000_000  # per item; beyond this the rest is not indexed (recorded as truncated)
MAX_PLAIN_BYTES = 8 * 1024 * 1024
CHUNK_TARGET = 1000
CHUNK_MAX = 1500
_WS = re.compile(r"[ \t\r\f\v]+")
_BLANKS = re.compile(r"\n{3,}")


@dataclass
class Section:
    """A run of text with an optional page number (PDF pages, spreadsheet sheets, slides)."""

    text: str
    page: int | None = None


@dataclass
class Extracted:
    sections: list[Section]
    truncated: bool = False
    extra: dict[str, Any] | None = None

    @property
    def chars(self) -> int:
        return sum(len(s.text) for s in self.sections)


class ExtractionError(RuntimeError):
    pass


class UnsupportedFormat(ExtractionError):
    """A known format there is no text extractor for yet: nothing to read, not a broken file."""


NO_EXTRACTOR = {".odt": "OpenDocument text", ".ods": "OpenDocument spreadsheet", ".odp": "OpenDocument slides"}


def clean(text: str) -> str:
    text = text.replace("\x00", "")
    lines = [_WS.sub(" ", line).strip() for line in text.split("\n")]
    return _BLANKS.sub("\n\n", "\n".join(lines)).strip()


_UTF16_BOMS = (b"\xff\xfe", b"\xfe\xff")


def _trim_cut_character(data: bytes) -> bytes:
    """Bytes read up to a size limit can end in the middle of a character: drop that partial character."""
    if data.startswith(_UTF16_BOMS):
        return data[: len(data) - len(data) % 2]
    try:
        data.decode("utf-8")
    except UnicodeDecodeError as exc:
        if exc.reason == "unexpected end of data":
            return data[: exc.start]
    return data


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        if encoding == "utf-16" and not data.startswith(_UTF16_BOMS):
            continue
        return text
    return data.decode("cp1252", errors="replace")


def looks_binary(data: bytes) -> bool:
    sample = data[:4096]
    if sample.startswith(_UTF16_BOMS):
        return False  # UTF-16 text: every other byte of ASCII is a zero, so the ratio below doesn't apply
    if b"\x00" in sample:
        return True
    printable = sum(32 <= b < 127 or b in (9, 10, 13) or b >= 128 for b in sample)
    return bool(sample) and printable / len(sample) < 0.85


class _HtmlText(HTMLParser):
    _SKIP: ClassVar[set[str]] = {"script", "style", "noscript", "template", "svg"}
    _BLOCK: ClassVar[set[str]] = {
        "p",
        "div",
        "br",
        "li",
        "tr",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "section",
        "article",
        "pre",
    }

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.title = ""
        self._skip = 0
        self._in_title = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self._SKIP:
            self._skip += 1
        elif tag == "title":
            self._in_title = True
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP and self._skip:
            self._skip -= 1
        elif tag == "title":
            self._in_title = False
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._in_title:
            self.title += data
        elif not self._skip:
            self.parts.append(data)


def _html(text: str) -> Extracted:
    parser = _HtmlText()
    parser.feed(text)
    return Extracted([Section(clean("".join(parser.parts)))], extra={"html_title": parser.title.strip()[:300]})


def _plain(path: Path, ext: str) -> Extracted:
    size = path.stat().st_size
    with path.open("rb") as fh:
        data = fh.read(MAX_PLAIN_BYTES)
    if looks_binary(data):
        raise ExtractionError("file looks binary")
    truncated = size > MAX_PLAIN_BYTES
    text = _decode(_trim_cut_character(data) if truncated else data)
    if ext in (".html", ".htm"):
        result = _html(text)
        result.truncated = truncated
        return result
    if ext in (".md", ".markdown"):
        front, body = parse_frontmatter(text)
        extra = {"frontmatter": {k: v for k, v in front.items() if isinstance(v, str | int | float | bool | list)}}
        return Extracted([Section(clean(body))], truncated=truncated, extra=extra)
    if ext in (".csv", ".tsv"):
        return _delimited(text, "\t" if ext == ".tsv" else ",", truncated)
    if ext == ".json":
        return Extracted([Section(_json_text(text))], truncated=truncated)
    return Extracted([Section(clean(text) if ext not in _CODE_KEEP_INDENT else text.replace("\x00", ""))], truncated)


_CODE_KEEP_INDENT = {".py", ".yaml", ".yml", ".toml"}


def _json_text(text: str) -> str:
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        return clean(text)
    strings: list[str] = []

    def walk(value: object, key: str = "") -> None:
        if len(strings) > 20_000:
            return
        if isinstance(value, dict):
            for k, v in value.items():
                walk(v, str(k))
        elif isinstance(value, list):
            for v in value:
                walk(v, key)
        elif isinstance(value, str | int | float) and not isinstance(value, bool):
            strings.append(f"{key}: {value}" if key else str(value))

    walk(data)
    return "\n".join(strings)


def _delimited(text: str, delimiter: str, truncated: bool) -> Extracted:
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = []
    header: list[str] | None = None
    for i, row in enumerate(reader):
        if i == 0:
            header = row
        rows.append(" | ".join(cell.strip() for cell in row if cell.strip()))
        if i > 50_000:
            truncated = True
            break
    return Extracted([Section("\n".join(rows))], truncated=truncated, extra={"columns": header or []})


def _pdf(path: Path) -> Extracted:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise ExtractionError("PDF is password-protected") from exc
        sections, total = [], 0
        for number, page in enumerate(reader.pages, start=1):
            text = clean(page.extract_text() or "")
            if text:
                sections.append(Section(text, page=number))
                total += len(text)
            if total > MAX_TEXT_CHARS:
                return Extracted(sections, truncated=True, extra={"pages": len(reader.pages)})
        info = reader.metadata or {}
        extra: dict[str, Any] = {"pages": len(reader.pages)}
        for key, name in (("/Title", "pdf_title"), ("/Author", "author"), ("/Subject", "subject")):
            value = info.get(key)
            if isinstance(value, str) and value.strip():
                extra[name] = value.strip()[:300]
        if not sections:
            extra["no_text_layer"] = True  # scanned PDF: OCR of PDF pages needs a page renderer (later phase)
        return Extracted(sections, extra=extra)
    except PdfReadError as exc:
        raise ExtractionError(f"unreadable PDF: {exc}") from exc


def _docx(path: Path) -> Extracted:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
    props = document.core_properties
    extra = {k: v for k, v in (("doc_title", props.title), ("author", props.author), ("subject", props.subject)) if v}
    return Extracted([Section(clean("\n".join(parts)))], extra=extra)


def _xlsx(path: Path) -> Extracted:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    sections, names, total, truncated = [], [], 0, False
    try:
        for index, sheet in enumerate(workbook.worksheets, start=1):
            names.append(sheet.title)
            lines = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if cells:
                    lines.append(" | ".join(cells))
                    total += sum(len(c) for c in cells)
                if total > MAX_TEXT_CHARS:
                    truncated = True
                    break
            if lines:
                sections.append(Section(f"{sheet.title}\n" + "\n".join(lines), page=index))
            if truncated:
                break
    finally:
        workbook.close()
    return Extracted(sections, truncated=truncated, extra={"sheets": names})


def _pptx(path: Path) -> Extracted:
    from pptx import Presentation

    deck = Presentation(str(path))
    sections = []
    for number, slide in enumerate(deck.slides, start=1):
        texts = []
        for shape in slide.shapes:
            frame = getattr(shape, "text_frame", None) if getattr(shape, "has_text_frame", False) else None
            if frame is not None and frame.text.strip():
                texts.append(frame.text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                texts.append(notes)
        if texts:
            sections.append(Section(clean("\n".join(texts)), page=number))
    return Extracted(sections, extra={"slides": len(deck.slides)})


def _url_file(path: Path) -> Extracted:
    text = _decode(path.read_bytes()[:65536])
    match = re.search(r"^URL=(.+)$", text, re.MULTILINE) or re.search(r"<string>(https?://[^<]+)</string>", text)
    url = match.group(1).strip() if match else ""
    return Extracted([Section(url)], extra={"url": url} if url else {})


def extract_text(path: Path) -> Extracted:
    """Extract text from a file. Raises ExtractionError for unreadable or binary files."""
    ext = path.suffix.lower()
    if ext in NO_EXTRACTOR:
        raise UnsupportedFormat(f"no text extractor for {NO_EXTRACTOR[ext]} ({ext}) files yet")
    try:
        if ext == ".pdf":
            result = _pdf(path)
        elif ext == ".docx":
            result = _docx(path)
        elif ext in (".xlsx", ".xlsm"):
            result = _xlsx(path)
        elif ext == ".pptx":
            result = _pptx(path)
        elif ext in (".url", ".webloc"):
            result = _url_file(path)
        elif ext in PLAIN_TEXT_EXTENSIONS or not ext:
            result = _plain(path, ext)
        else:
            result = _plain(path, ext)  # unknown extension: index it if it's actually text
    except ExtractionError:
        raise
    except (OSError, ValueError, KeyError, IndexError, TypeError, AttributeError, ImportError) as exc:
        raise ExtractionError(f"{type(exc).__name__}: {exc}"[:500]) from exc
    total = 0
    kept = []
    for section in result.sections:
        if total + len(section.text) > MAX_TEXT_CHARS:
            section = Section(section.text[: MAX_TEXT_CHARS - total], section.page)
            result.truncated = True
        kept.append(section)
        total += len(section.text)
        if total >= MAX_TEXT_CHARS:
            break
    result.sections = [s for s in kept if s.text.strip()]
    return result


# ---------- chunking ----------


@dataclass(frozen=True)
class Chunk:
    text: str
    page: int | None
    char_start: int
    char_end: int


def _pieces(text: str) -> Iterator[tuple[int, int]]:
    """Paragraph spans; paragraphs longer than CHUNK_MAX are cut at sentence or word boundaries."""
    for match in re.finditer(r"[^\n]+(?:\n(?!\n)[^\n]+)*", text):
        start, end = match.span()
        while end - start > CHUNK_MAX:
            window = text[start : start + CHUNK_MAX]
            cut = max(window.rfind(". "), window.rfind("? "), window.rfind("! "), window.rfind("\n"))
            if cut < CHUNK_TARGET // 2:
                cut = window.rfind(" ")
            if cut < CHUNK_TARGET // 2:
                cut = CHUNK_MAX - 1
            yield start, start + cut + 1
            start += cut + 1
            while start < end and text[start].isspace():
                start += 1
        if end > start:
            yield start, end


def chunk_sections(sections: list[Section]) -> list[Chunk]:
    """Pack paragraphs into ~1000-character chunks, never crossing a page."""
    chunks: list[Chunk] = []
    for section in sections:
        text = section.text
        current_start: int | None = None
        current_end = 0
        for start, end in _pieces(text):
            if current_start is not None and end - current_start > CHUNK_TARGET:
                chunks.append(Chunk(text[current_start:current_end].strip(), section.page, current_start, current_end))
                current_start = None
            if current_start is None:
                current_start = start
            current_end = end
        if current_start is not None and text[current_start:current_end].strip():
            chunks.append(Chunk(text[current_start:current_end].strip(), section.page, current_start, current_end))
    return chunks
