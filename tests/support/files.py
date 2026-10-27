"""Real files of every core media type, generated locally (no downloads), for tests and the seed generator.

Each writer produces a genuine file the production extractors read: a hand-built PDF with a text layer,
DOCX/XLSX/PPTX through their libraries, JPEGs with EXIF (date, camera, GPS), PNGs with rendered text for
OCR, and WAVs with RIFF INFO tags.
"""

import struct
import wave
from pathlib import Path
from typing import Any

from PIL import ExifTags, Image, ImageDraw, ImageFont


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def write_pdf(path: Path, pages: list[str], *, title: str | None = None) -> Path:
    """A minimal valid PDF 1.4: one Helvetica text block per page, with a real xref table."""
    objects: list[bytes] = []
    kids = []
    font_id = 3 + 2 * len(pages)
    for index, text in enumerate(pages):
        page_id, content_id = 3 + 2 * index, 4 + 2 * index
        kids.append(f"{page_id} 0 R")
        lines = text.split("\n")
        ops = ["BT", "/F1 12 Tf", "14 TL", "72 760 Td"] + [f"({_pdf_escape(line)}) Tj T*" for line in lines] + ["ET"]
        stream = "\n".join(ops).encode("latin-1", errors="replace")
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 {font_id} 0 R >> >>"
            f" /Contents {content_id} 0 R >>".encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    info_id = font_id + 1
    objects.append(f"<< /Title ({_pdf_escape(title or path.stem)}) /Producer (UNIVERSE tests) >>".encode())
    header = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>".encode(),
    ]
    body = header + objects
    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for number, obj in enumerate(body, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(body) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(body) + 1} /Root 1 0 R /Info {info_id} 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))
    return path


def write_docx(path: Path, title: str, paragraphs: list[str]) -> Path:
    import docx

    document = docx.Document()
    document.core_properties.title = title
    document.add_heading(title, level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(str(path))
    return path


def write_xlsx(path: Path, sheets: dict[str, list[list[object]]]) -> Path:
    from openpyxl import Workbook

    workbook = Workbook()
    first = True
    for name, rows in sheets.items():
        sheet = workbook.active if first else workbook.create_sheet()
        assert sheet is not None
        sheet.title = name
        for row in rows:
            sheet.append(row)
        first = False
    workbook.save(str(path))
    return path


def write_pptx(path: Path, slides: list[tuple[str, str]]) -> Path:
    from pptx import Presentation

    deck = Presentation()
    for heading, body in slides:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        title: Any = slide.shapes.title
        placeholder: Any = slide.placeholders[1]
        title.text = heading
        placeholder.text = body
    deck.save(str(path))
    return path


def write_text_image(path: Path, text: str, *, size: tuple[int, int] = (900, 200)) -> Path:
    """Black text on white, large enough for OCR to read reliably."""
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    draw.text((20, size[1] // 2 - 30), text, fill="black", font=ImageFont.load_default(size=44))
    image.save(path)
    return path


def write_photo(
    path: Path,
    *,
    color: tuple[int, int, int] = (40, 90, 160),
    taken: str | None = "2024:07:14 19:30:00",
    camera: tuple[str, str] | None = ("Canon", "EOS R6"),
    gps: tuple[float, float] | None = None,
    size: tuple[int, int] = (640, 480),
) -> Path:
    image = Image.new("RGB", size, color)
    draw = ImageDraw.Draw(image)
    draw.ellipse((size[0] // 3, size[1] // 3, size[0] // 2, size[1] // 2), fill=(250, 200, 80))
    exif = Image.Exif()
    if camera:
        exif[ExifTags.Base.Make], exif[ExifTags.Base.Model] = camera
    if taken:
        exif[ExifTags.Base.DateTime] = taken
        exif.get_ifd(ExifTags.IFD.Exif)[ExifTags.Base.DateTimeOriginal] = taken
    if gps:

        def dms(value: float) -> tuple[float, float, float]:
            value = abs(value)
            d = int(value)
            m = int((value - d) * 60)
            return (float(d), float(m), round((value - d - m / 60) * 3600, 4))

        gps_ifd = exif.get_ifd(ExifTags.IFD.GPSInfo)
        gps_ifd[1] = "N" if gps[0] >= 0 else "S"
        gps_ifd[2] = dms(gps[0])
        gps_ifd[3] = "E" if gps[1] >= 0 else "W"
        gps_ifd[4] = dms(gps[1])
    image.save(path, "JPEG", exif=exif, quality=85)
    return path


def write_wav(path: Path, *, title: str, artist: str, album: str = "", seconds: float = 0.5) -> Path:
    """Silent PCM WAV with a RIFF LIST/INFO chunk (title, artist, album) that tag readers understand."""
    rate = 8000
    frames = b"\x00\x00" * int(rate * seconds)
    with wave.open(str(path), "wb") as out:
        out.setnchannels(1)
        out.setsampwidth(2)
        out.setframerate(rate)
        out.writeframes(frames)

    def sub(tag: bytes, value: str) -> bytes:
        data = value.encode("utf-8") + b"\x00"
        if len(data) % 2:
            data += b"\x00"
        return tag + struct.pack("<I", len(data)) + data

    info = b"INFO" + sub(b"INAM", title) + sub(b"IART", artist) + (sub(b"IPRD", album) if album else b"")
    raw = bytearray(path.read_bytes())
    raw += b"LIST" + struct.pack("<I", len(info)) + info
    raw[4:8] = struct.pack("<I", len(raw) - 8)
    path.write_bytes(bytes(raw))
    return path
